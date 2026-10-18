"""Build the VP pitch deck: black background, white text, Cognition logo top right."""
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt

WHITE, GREY, DIM, BLACK = (RGBColor(0xFF, 0xFF, 0xFF), RGBColor(0x9A, 0x9A, 0x9A),
                           RGBColor(0x55, 0x55, 0x55), RGBColor(0, 0, 0))
FONT = "Arial"
W, H = Inches(13.333), Inches(7.5)
L = Inches(0.9)

prs = Presentation()
prs.slide_width, prs.slide_height = W, H
blank = prs.slide_layouts[6]
n = 0


def text(slide, x, y, w, h, runs, size=22, color=WHITE, bold=False, space=10):
    """runs: list of paragraphs; each paragraph is a str or (str, size, color, bold)."""
    tf = slide.shapes.add_textbox(x, y, w, h).text_frame
    tf.word_wrap = True
    for i, r in enumerate(runs):
        t, s, c, b = (r, size, color, bold) if isinstance(r, str) else r
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(space)
        run = p.add_run()
        run.text = t
        run.font.name, run.font.size, run.font.color.rgb, run.font.bold = FONT, Pt(s), c, b
    return tf


def slide(kicker=None, title=None):
    global n
    n += 1
    s = prs.slides.add_slide(blank)
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = BLACK
    s.shapes.add_picture("logo_white.png", W - Inches(2.35), Inches(0.45), height=Inches(0.32))
    if n > 1:
        text(s, W - Inches(1.4), H - Inches(0.7), Inches(0.8), Inches(0.4), [(str(n), 12, DIM, False)])
    if kicker:
        text(s, L, Inches(0.95), Inches(10), Inches(0.5), [(kicker.upper(), 13, GREY, True)])
    if title:
        text(s, L, Inches(1.35), Inches(11.5), Inches(1.4), [(title, 40, WHITE, True)])
    return s


def columns(s, cols, y=Inches(3.2), size=20, head=24):
    gap = Inches(0.4)
    cw = int((W - 2 * L - gap * (len(cols) - 1)) / len(cols))
    for i, (h, lines) in enumerate(cols):
        x = L + i * (cw + gap)
        rule = s.shapes.add_shape(1, x, y - Inches(0.15), cw, Pt(1.5))
        rule.fill.solid()
        rule.fill.fore_color.rgb = DIM
        rule.line.fill.background()
        text(s, x, y, cw, Inches(3.5), [(h, head, WHITE, True)] + [(t, size, GREY, False) for t in lines])


def stats(s, items, y=Inches(3.0)):
    gap = Inches(0.4)
    cw = int((W - 2 * L - gap * (len(items) - 1)) / len(items))
    for i, (big, label) in enumerate(items):
        x = L + i * (cw + gap)
        text(s, x, y, cw, Inches(2.6), [(big, 60, WHITE, True), (label, 18, GREY, False)], space=6)


def footer(s, t):
    text(s, L, H - Inches(1.45), Inches(11.3), Inches(0.8), [(t, 18, GREY, False)])


# 1. Title
s = slide()
text(s, L, Inches(2.5), Inches(11), Inches(3), [
    ("Power Apps or Devin?", 54, WHITE, True),
    ("Internal tools for a 60-engineer fintech", 26, GREY, False),
])
text(s, L, H - Inches(1.4), Inches(8), Inches(0.5), [("For the VP of Engineering", 16, DIM, False)])

# 2. The answer
s = slide("The answer", "Build with Devin, as a gated migration.")
text(s, L, Inches(3.3), Inches(11), Inches(2.5), [
    ("The decision isn't the 3 tools you have.", 26, GREY, False),
    ("It's the 10+ tools you haven't specified yet.", 26, WHITE, True),
])

# 3. What Power Apps sells (brief point 1)
s = slide("What Power Apps really sells", "Speed × governance, repeated.")
columns(s, [
    ("Speed", ["Building blocks ship a tool in days"]),
    ("Governance", ["SSO, roles, masking, audit, DLP: all configuration"]),
    ("Managed platform", ["Microsoft hosts, patches and certifies it"]),
])
footer(s, "Where it runs out:  custom logic  ·  your own APIs  ·  automated testing  ·  scale")

# 4. How Devin helps: the platform (brief point 2)
s = slide("How Devin helps: what I built", "Built once. Inherited by every tool.")
columns(s, [
    ("Real SSO", ["Roles come from identity-provider groups"]),
    ("Masked PII", ["Reveal needs a reason, and every view is audited"]),
    ("Approvals", ["No self-approval; 2 approvers over $1K"]),
    ("Audit", ["Append-only, written with every change"]),
], size=18, head=22)
footer(s, "Prototype: one shared Python platform + rules file (AGENTS.md). Built in ~47 minutes.")

# 5. Refunds
s = slide("Beyond the building blocks", "Your riskiest workflow, not a form.")
text(s, L, Inches(3.2), Inches(11.5), Inches(2.5), [
    ("Provider times out  →  UNKNOWN  →  reconcile  →  never pays twice", 24, WHITE, True),
    ("Every business rule is a named test that runs on every PR.", 22, GREY, False),
], space=24)

# 6. Tool factory
s = slide("The tool factory", "One paragraph in. A governed tool out.")
stats(s, [
    ("6 / 8 min", "Fresh Devin sessions,\nprompt → PR with green CI"),
    ("0", "Code bugs found\nin review"),
    ("1", "Review round each\nbefore merge"),
])
footer(s, "Controls were inherited without being asked for. Review caught decisions nobody wrote down; that's the control point.")

# 7. Honest assessment (brief point 3)
s = slide("Honest assessment", "What's replicable, and what isn't.")
columns(s, [
    ("Replicable with Devin", ["SSO and roles", "Masking and audited reveal", "Audit and approvals",
                               "Custom logic and your APIs"]),
    ("Not replicable", ["Ops staff editing tools themselves", "Microsoft's hosting and certifications", "Mobile / offline"]),
])

# 8. Risk
s = slide("Where the risk lies", "You own the security surface.")
columns(s, [
    ("AI-written code", ["Touches money and PII", "Human review and rule tests are mandatory"]),
    ("Compliance", ["Pen test, encryption at rest", "AML retention vs GDPR erasure"]),
    ("New vendor", ["Devin has code access", "Vendor review; no real PII in sessions"]),
])
footer(s, "Power Apps doesn't remove this either: Microsoft certifies the platform, not your apps.")

# 9. Cost (brief point 4)
s = slide("3-year cost", "Build is cheaper over 3 years.")
stats(s, [
    ("$1.20M", "Buy\nPower Apps"),
    ("$1.04M", "Build\nbase case"),
    ("$0.79M", "Build\nDevin does the upkeep"),
])
footer(s, "Year 1 costs ~$220K more to build. Build overtakes Buy around month 26.")

# 10. Engineering capacity
s = slide("The metric that matters", "Engineers needed, from year 2.")
stats(s, [
    ("0.52", "FTE\nBuy"),
    ("0.71", "FTE\nBuild, half-time owner"),
    ("0.36", "FTE\nBuild, Devin does the upkeep"),
])
footer(s, "Base case favours Buy. If the owner only reviews Devin's PRs (~6 h/week), Build needs fewer engineers.")

# 11. Recommendation
s = slide("Recommendation", "Build, gated on real numbers.")
columns(s, [
    ("Why build", ["Your next tools are custom", "The licence scales with users", "Devin cuts the engineers needed"]),
    ("What flips it to buy", ["Next 10 tools are simple forms", "Licence under ~$185K",
                              "Ops staff build tools today", "Nobody can own the platform"]),
])

# 12. Next steps (brief point 5)
s = slide("Next steps", "Prove it on your hardest tool.")
columns(s, [
    ("Week 0", ["Licence breakdown", "List of the next 10 tools", "Card data? EU-regulated?"]),
    ("Weeks 1–6", ["Devin + 1 engineer", "Harden the platform", "Highest-value tool first, on real systems"]),
    ("Gate", ["Human hours per tool ≤ 20% of Power Apps estimate", "Owner ≤ 0.25 FTE", "No high-severity failures"]),
], size=18, head=22)
footer(s, "Drop Power Apps at renewal only if the gate passes.")

# 13. If they buy
s = slide("If you keep Power Apps", "Devin still fits.")
columns(s, [
    ("Optimise the licence", ["Per-app or pay-as-you-go for light users", "Up to ~$100K a year saved"]),
    ("Use Devin at the edges", ["Custom connectors, code components, Azure Functions", "Where code is needed anyway"]),
])

prs.save("Power_Apps_vs_Devin.pptx")
print("slides", n)
